"""CS-114: the toxicity weights E3 multiplies TRI poundage by.

Two things here are worth more than the rest. The first is that a chemical RSEI
cannot weight produces no row at all rather than a row weighted zero, because a
zero would say "harmless" about a chemical EPA has merely never assessed for
inhalation, and E3 would then rank a facility releasing it as cleanly as one
releasing nothing. The second is the CAS reconstruction: the column that would
be the obvious join key has been corrupted by Excel in EPA's published file, and
the test that pins that is the only thing standing between this adapter and
silently dropping the largest unweighted release in the pilot state.
"""

import io

import httpx
import pytest
from openpyxl import Workbook

from pipeline.adapters.base import FetchResult
from pipeline.adapters.rsei import (
    PINNED_EDITION,
    ChemicalToxicityWeight,
    EpaRseiAdapter,
    RseiChemical,
    check_digit_ok,
    standard_cas,
)
from pipeline.errors import PermanentSourceError, RecordRejected
from tests.conftest import make_context, make_fetcher

# The columns the adapter reads, in the order the published workbook lists them.
HEADER = ("CASNumber", "CASStandard", "Chemical", "ITW", "ToxicityClassInhale")

# Benzene, as v2312 publishes it.
BENZENE = ("71432", "71-43-2", "Benzene", 28000, "Cancer")

# Hydrogen sulfide, as v2312 publishes it: the packed CAS is intact and the
# hyphenated one has been turned into an Excel date serial. See the module
# docstring of pipeline/adapters/rsei.py.
HYDROGEN_SULFIDE = ("7783064", "2148878", "Hydrogen sulfide", 1800, "Non-cancer")

# A chemical category. TRI reports these under an N code and so does RSEI.
ARSENIC = ("N020", "N020", "Arsenic compounds", 15000000, "Cancer")

# Allylamine: on the TRI list, with no inhalation toxicity data in RSEI.
UNWEIGHTED = ("107119", "107-11-9", "Allylamine", None, None)


def workbook_bytes(
    rows: tuple[tuple[object, ...], ...] = (BENZENE,),
    *,
    version: str | None = "Version 2.3.12",
    header: tuple[str, ...] = HEADER,
    toxicity_sheet: str = "Toxicity data",
) -> bytes:
    """A workbook shaped like EPA's, built from the rows a test cares about."""
    book = Workbook()
    description = book.active
    description.title = "Description"
    description.append(["DATE:March  2024"])
    description.append(["This spreadsheet contains the toxicity information."])
    if version is not None:
        description.append([version])

    sheet = book.create_sheet(toxicity_sheet)
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def written_weight(adapter: EpaRseiAdapter, record: RseiChemical) -> ChemicalToxicityWeight:
    """`normalize` is typed to the base record; these tests assert on the subclass."""
    (written,) = list(adapter.normalize(record, ctx=None))  # type: ignore[arg-type]
    assert isinstance(written, ChemicalToxicityWeight)
    return written


def chemical(**overrides: object) -> RseiChemical:
    fields: dict[str, object] = {
        "cas_number": "71432",
        "reported_cas_standard": "71-43-2",
        "chemical_name": "Benzene",
        "weight": 28000.0,
        "toxicity_class": "Cancer",
    }
    fields.update(overrides)
    return RseiChemical(**fields)  # type: ignore[arg-type]


async def fetch_from(payload: bytes) -> FetchResult[RseiChemical]:
    """Run `fetch` against a transport serving one workbook."""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=payload)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        adapter = EpaRseiAdapter()
        http = make_fetcher(client, source="epa_rsei")
        ctx = make_context(http=http, sink=None, source="epa_rsei")  # type: ignore[arg-type]
        return await adapter.fetch(ctx)


# ---- the CAS number ---------------------------------------------------


def test_a_published_cas_number_verifies_against_its_own_check_digit() -> None:
    """Benzene, acetaldehyde and hydrogen sulfide, digit for digit."""
    assert check_digit_ok("71432")
    assert check_digit_ok("75070")
    assert check_digit_ok("7783064")


def test_a_transposed_cas_number_fails_its_check_digit() -> None:
    """The check exists to catch a cell that was altered on its way out of Excel."""
    assert not check_digit_ok("71423")
    assert not check_digit_ok("2148878")


def test_something_too_short_to_be_a_cas_number_is_not_one() -> None:
    assert not check_digit_ok("714")
    assert not check_digit_ok("N020")


def test_the_hyphenated_form_is_rebuilt_from_the_packed_one() -> None:
    """CAS format is n-nn-n: a check digit, two digits, and the rest in front."""
    assert standard_cas("71432") == "71-43-2"
    assert standard_cas("7783064") == "7783-06-4"
    assert standard_cas("1918021") == "1918-02-1"


async def test_a_cas_excel_turned_into_a_date_still_joins_to_tri() -> None:
    """The regression this adapter exists around.

    EPA's own hyphenated column reads 2148878 for hydrogen sulfide, an Excel date
    serial for the year 7783. Joining on it would leave the largest unweighted
    release in Louisiana's 2024 extract, 930,064 lb, outside E3 while looking
    like a chemical RSEI had simply never assessed.
    """
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"
    record = chemical(
        cas_number="7783064", reported_cas_standard="2148878", chemical_name="Hydrogen sulfide"
    )

    adapter.validate(record, ctx=None)  # type: ignore[arg-type]

    assert written_weight(adapter, record).cas_number == "7783-06-4"


def test_a_reconstruction_that_contradicts_an_intact_column_is_rejected() -> None:
    """Where EPA's hyphenated cell survived, it is the check on the rebuild."""
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="disagrees"):
        adapter.validate(
            chemical(cas_number="71432", reported_cas_standard="71-43-9"),
            ctx=None,  # type: ignore[arg-type]
        )


def test_a_chemical_category_keeps_its_n_code_untouched() -> None:
    """`tri_release.cas_number` holds N020 for arsenic compounds, not a CAS number."""
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"
    record = chemical(
        cas_number="N020",
        reported_cas_standard="N020",
        chemical_name="Arsenic compounds",
        weight=15_000_000.0,
    )

    adapter.validate(record, ctx=None)  # type: ignore[arg-type]

    assert written_weight(adapter, record).cas_number == "N020"


# ---- validation -------------------------------------------------------


def test_a_row_with_no_cas_number_is_rejected() -> None:
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="missing CAS"):
        adapter.validate(chemical(cas_number=""), ctx=None)  # type: ignore[arg-type]


def test_a_tri_placeholder_is_skipped_rather_than_rejected() -> None:
    """RSEI carries INVALID, MIXTURE and TRD SECRT so its table lines up with TRI's.

    Skipped and counted rather than rejected, on the same grounds the TRI adapter
    holds a Form A filing as absent: there are exactly three, they are there every
    edition, and rejecting them would mark every healthy pull `partial` for a
    condition nobody should ever go and look at.
    """
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"
    for placeholder in ("MIXTURE", "TRD SECRT", "INVALID"):
        record = chemical(cas_number=placeholder, reported_cas_standard=placeholder, weight=None)
        adapter.validate(record, ctx=None)  # type: ignore[arg-type]
        assert list(adapter.normalize(record, ctx=None)) == []  # type: ignore[arg-type]


def test_a_placeholder_is_never_weighted_even_if_upstream_gives_it_a_number() -> None:
    """Its identifier could never join `tri_release.cas_number`, so a row would be junk."""
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"
    record = chemical(cas_number="MIXTURE", reported_cas_standard="MIXTURE", weight=42.0)

    adapter.validate(record, ctx=None)  # type: ignore[arg-type]
    assert list(adapter.normalize(record, ctx=None)) == []  # type: ignore[arg-type]


async def test_a_healthy_pull_rejects_nothing() -> None:
    """The rejection count is the signal that something changed upstream.

    Every row of a good edition is either written or deliberately skipped, so a
    rejection here means a CAS number, a weight or a column stopped looking like
    itself.
    """
    placeholder = ("MIXTURE", "MIXTURE", "MIXTURE", None, None)
    result = await fetch_from(workbook_bytes(rows=(BENZENE, UNWEIGHTED, placeholder)))

    adapter = EpaRseiAdapter()
    for record in result.records:
        adapter.validate(record, ctx=None)  # type: ignore[arg-type]
    assert any("1 rows are TRI placeholders" in note for note in result.notes)


def test_a_cas_number_that_fails_its_check_digit_is_rejected_and_counted() -> None:
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="check digit"):
        adapter.validate(
            chemical(cas_number="71423", reported_cas_standard=""),
            ctx=None,  # type: ignore[arg-type]
        )


def test_a_weight_of_zero_is_rejected_rather_than_stored() -> None:
    """Zero is the one value that must never reach this table.

    The schema forbids it, and the reason the schema forbids it is that a zero
    weight and an absent weight mean opposite things: one says the chemical is
    harmless, the other says EPA never assessed it.
    """
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="not positive"):
        adapter.validate(chemical(weight=0.0), ctx=None)  # type: ignore[arg-type]


def test_a_negative_weight_is_rejected() -> None:
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="not positive"):
        adapter.validate(chemical(weight=-1.0), ctx=None)  # type: ignore[arg-type]


def test_a_weight_cell_that_is_not_a_number_is_rejected_not_read_as_absent() -> None:
    """A blank means no data. Text in the cell means the column moved."""
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="not positive"):
        adapter.validate(chemical(weight=float("nan")), ctx=None)  # type: ignore[arg-type]


def test_a_row_with_no_chemical_name_is_rejected() -> None:
    adapter = EpaRseiAdapter()
    with pytest.raises(RecordRejected, match="missing chemical name"):
        adapter.validate(chemical(chemical_name=""), ctx=None)  # type: ignore[arg-type]


# ---- normalize --------------------------------------------------------


def test_a_chemical_rsei_cannot_weight_produces_no_row_at_all() -> None:
    """Acceptance criterion two.

    362 of v2312's 823 chemicals are in this state. Writing them as zeros would
    roughly double the table and quietly declare every one of them harmless; the
    E3 join finds no row instead and leaves their poundage out of the indicator.
    """
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"
    record = chemical(
        cas_number="107119",
        reported_cas_standard="107-11-9",
        chemical_name="Allylamine",
        weight=None,
        toxicity_class="",
    )

    adapter.validate(record, ctx=None)  # type: ignore[arg-type]
    assert list(adapter.normalize(record, ctx=None)) == []  # type: ignore[arg-type]


def test_every_written_row_carries_the_edition_it_came_from() -> None:
    """So the weight a facility's contribution was computed with stays recoverable."""
    adapter = EpaRseiAdapter()
    adapter._edition = "v2312"

    written = written_weight(adapter, chemical())

    assert written.source_edition == "v2312"
    assert written.table == "chemical_toxicity_weight"


def test_the_natural_key_is_the_cas_number_alone() -> None:
    """One weight per chemical, so a re-pull replaces rather than duplicates."""
    written = ChemicalToxicityWeight(
        cas_number="71-43-2", chemical_name="Benzene", rsei_weight=28000.0, source_edition="v2312"
    )
    assert written.natural_key() == ("71-43-2",)


def test_an_oral_adopted_weight_is_recognised_by_its_asterisk() -> None:
    """EPA's own convention, and 238 of the 461 weights carry it."""
    assert chemical(toxicity_class="Non-cancer*").oral_adopted
    assert not chemical(toxicity_class="Cancer").oral_adopted


# ---- fetch ------------------------------------------------------------


async def test_the_vintage_is_the_rsei_version_the_workbook_declares() -> None:
    """Acceptance criterion one. Not a year, and not the night of the download."""
    result = await fetch_from(workbook_bytes())
    assert result.vintage == PINNED_EDITION == "v2312"


async def test_the_download_is_checksummed_into_the_manifest() -> None:
    result = await fetch_from(workbook_bytes())
    (artifact,) = result.artifacts
    assert artifact.sha256 and artifact.url.endswith(".xlsx")


async def test_a_replaced_edition_is_labelled_by_the_file_not_by_the_pin() -> None:
    """EPA replacing the file in place must not load new weights under an old name."""
    result = await fetch_from(workbook_bytes(version="Version 2.4.1"))
    assert result.vintage == "v241"
    assert any("replaced the file" in note for note in result.notes)


async def test_a_workbook_that_declares_no_version_fails_the_pull() -> None:
    """An unlabelled weight cannot be told from the edition it replaced."""
    with pytest.raises(PermanentSourceError, match="no RSEI version"):
        await fetch_from(workbook_bytes(version=None))


async def test_a_renamed_column_fails_the_pull_by_name() -> None:
    """Rather than loading a table of empty weights."""
    header = ("CASNumber", "CASStandard", "Chemical", "InhalationTW", "ToxicityClassInhale")
    with pytest.raises(PermanentSourceError, match="ITW"):
        await fetch_from(workbook_bytes(header=header))


async def test_a_renamed_sheet_fails_the_pull() -> None:
    with pytest.raises(PermanentSourceError, match="Toxicity data"):
        await fetch_from(workbook_bytes(toxicity_sheet="Weights"))


async def test_something_that_is_not_a_workbook_fails_the_pull() -> None:
    with pytest.raises(PermanentSourceError, match="not a readable workbook"):
        await fetch_from(b"<html>503 Service Unavailable</html>")


async def test_a_sheet_with_only_a_header_fails_the_pull() -> None:
    with pytest.raises(PermanentSourceError, match="no data rows"):
        await fetch_from(workbook_bytes(rows=()))


async def test_two_weights_for_one_chemical_fail_the_pull() -> None:
    """The table is keyed on the CAS number, so one of the two would vanish.

    Which one survived would depend on row order, and the difference multiplies
    straight into E3.
    """
    twice = (BENZENE, ("71432", "71-43-2", "Benzene", 99, "Cancer"))
    with pytest.raises(PermanentSourceError, match="repeats"):
        await fetch_from(workbook_bytes(rows=twice))


async def test_a_pull_reports_how_many_chemicals_it_could_not_weight() -> None:
    """Acceptance criterion two's other half: excluded, and counted."""
    result = await fetch_from(workbook_bytes(rows=(BENZENE, HYDROGEN_SULFIDE, ARSENIC, UNWEIGHTED)))
    assert result.records[3].weight is None
    assert any("4 chemicals in RSEI" in note and "3 with an" in note for note in result.notes)


async def test_the_pull_reads_every_row_of_a_realistic_sheet() -> None:
    rows = (BENZENE, HYDROGEN_SULFIDE, ARSENIC, UNWEIGHTED)
    result = await fetch_from(workbook_bytes(rows=rows))

    assert [c.cas_number for c in result.records] == ["71432", "7783064", "N020", "107119"]
    assert [c.weighted for c in result.records] == [True, True, True, False]


# ---- what the source always has to declare ----------------------------


def test_the_source_declares_the_gaps_that_are_true_every_night() -> None:
    adapter = EpaRseiAdapter()
    gaps = adapter.known_gaps(ctx=None)  # type: ignore[arg-type]

    assert all(gap.affects == ("E3",) for gap in gaps)
    details = " ".join(gap.detail for gap in gaps)
    # The three a reader of docs/provenance.md most needs: what is unweighted,
    # that half the weights are not inhalation data, and that the edition is pinned.
    assert "461 of the 823" in details
    assert "adopted from the oral pathway" in details
    assert PINNED_EDITION in details


def test_the_adapter_declares_what_a_good_load_looks_like() -> None:
    """A source with no expectations is reported by the gate as a gap."""
    expectations = EpaRseiAdapter.expectations
    assert expectations is not None
    (table,) = expectations.tables
    assert table.table == "chemical_toxicity_weight"
    assert table.rows is not None and table.rows.high is not None
    assert table.rows.low < 461 < table.rows.high
